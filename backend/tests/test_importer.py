from decimal import Decimal
from backend.importer import calculate_total, import_transactions, import_income
from datetime import date
from openpyxl import Workbook
import pytest


def test_calculate_total():
    transactions = [
        {"amount":Decimal("1.45")},
        {"amount":Decimal("2.55")},
        {"amount":Decimal("3.00")}
    ]
    result = calculate_total(transactions)
    assert result == Decimal("7.00")

def test_calculate_total_empty():
    transactions = []
    result = calculate_total(transactions)
    assert result == Decimal("0.00")

def test_import_transaction(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Expense2026"
    worksheet["A1"] = "Month"
    worksheet["B1"] = "Grocery"
    worksheet["C1"] = "Gas"

    worksheet["A3"] = date(2026, 1, 1)
    worksheet["B3"] = 45.50
    worksheet["C3"] = 20.25  
    file_path = tmp_path / "test_expenses.xlsx"  
    workbook.save(file_path)

    transactions = import_transactions(file_path, "Expense2026", 3 , 3)

    assert len(transactions) == 2
    assert calculate_total(transactions) == Decimal("65.75")

def test_import_invalid_amount(tmp_path):   
     workbook = Workbook()
     worksheet = workbook.active
     worksheet.title = "Expense2026"
     worksheet["A1"] = "Month"
     worksheet["B1"] = "Grocery"
     worksheet["A3"] = date(2026, 1, 1)
     worksheet["B3"] = "hello"
     file_path = tmp_path / "test_expenses.xlsx"
     workbook.save(file_path)

     with pytest.raises(ValueError, match = "Invalid amount"):
         import_transactions(file_path, "Expense2026", 3, 3)

def test_missing_formula_result(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Expense2026"

    worksheet["A1"] = "Month"
    worksheet["B1"] = "Grocery"

    worksheet["A3"] = date(2026, 1, 1)
    worksheet["B3"] = "=10+5"

    file_path = tmp_path / "missing_formula.xlsx"
    workbook.save(file_path)

    with pytest.raises(ValueError, match="Missing calculated value in cell B3"):
        import_transactions(file_path, "Expense2026", 3, 3)

def test_import_income(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Expense2026"

    worksheet["A3"] = date(2026, 1, 1)
    worksheet["O3"] = 1000.00

    worksheet["A4"] = date(2026,1,15)
    worksheet["O4"] = 2000.00

    file_path = tmp_path / "test_income.xlsx"
    workbook.save(file_path)

    incomes = import_income(file_path, "Expense2026", 3, 4)

    assert len(incomes) == 2
    assert calculate_total(incomes) == Decimal("3000.00")
