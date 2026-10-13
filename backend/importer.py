from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal
from datetime import date, datetime


def calculate_total(transactions) -> Decimal:
    total = Decimal("0.00")
    for transaction in transactions:
        total += transaction["amount"]
    return total 


def import_transactions(file_path, sheet_name, min_row, max_row):
    workbook = load_workbook(file_path, data_only = True)
    formula_workbook = load_workbook(file_path, data_only = False)
    formula_worksheet = formula_workbook[sheet_name]
    worksheet = workbook[sheet_name]
    headers = [cell.value for cell in worksheet[1]]
    transactions = []    

    for row_number, row in enumerate(worksheet.iter_rows(min_row=min_row, max_row=max_row, max_col=13, values_only=True), start = min_row):
        excel_date = row[0]   
        if not isinstance(excel_date, date):
            raise ValueError(f"Invalid transaction date: {excel_date!r}")        
        transaction_date = excel_date.date() if isinstance(excel_date, datetime) else excel_date

        for i in range(1, 13):
            amount = row[i]           
            if amount is None:
                formula_cell = formula_worksheet.cell(row=row_number, column = i + 1)

                if formula_cell.data_type == "f":
                    raise ValueError(f"Missing calculated value in cell {formula_cell.coordinate}")
            if amount is not None:
                category = headers[i]                 
                if not isinstance(category,str) or not category.strip():
                    raise ValueError (f"Invalid expense category: {category!r}")
                if not isinstance(amount,(int, float)) or isinstance(amount, bool):
                    raise ValueError(f"Invalid amount in category {category}: {amount!r}")               
                transactions.append({
                    "date" : transaction_date,
                    "category": category,
                    "amount" : Decimal(str(amount)).quantize(Decimal("0.01")),
                    "source_cell": worksheet.cell(row=row_number, column=i+1).coordinate
                }) 
    return transactions  


def import_income(file_path, sheet_name, min_row, max_row):
    workbook = load_workbook(file_path, data_only=True)
    formula_workbook = load_workbook(file_path, data_only=False)

    worksheet = workbook[sheet_name]
    original_category = worksheet.cell(row=1, column=15).value
    formula_worksheet = formula_workbook[sheet_name]

    income_transactions = []

    for row_number in range(min_row, max_row + 1):
        amount = worksheet.cell(row=row_number, column=15).value

        if amount is None:
            formula_cell = formula_worksheet.cell(row=row_number, column=15)

            if formula_cell.data_type == "f":
                raise ValueError(
                    f"Missing calculated value in cell {formula_cell.coordinate}"
                )
            continue

        if not isinstance(amount, (int, float)) or isinstance(amount, bool):
            raise ValueError(f"Invalid income amount in row {row_number}: {amount!r}")

        excel_date = worksheet.cell(row=row_number, column=1).value

        if not isinstance(excel_date, date):
            raise ValueError(f"Invalid income date: {excel_date!r}")

        transaction_date = excel_date.date() if isinstance(excel_date, datetime) else excel_date

        income_transactions.append({
            "date": transaction_date,
            "category": "Income",
            "amount": Decimal(str(amount)).quantize(Decimal("0.01")),
            "source_cell": worksheet.cell(row=row_number, column=15).coordinate,
            "original_category": original_category
        })

    return income_transactions   


def reconcile_totals(file_path, sheet_name, summary_row, total_income, total, actual_savings):
    workbook = load_workbook(file_path, data_only=True)   
    worksheet = workbook[sheet_name]

    excel_income = Decimal(str(worksheet[f"O{summary_row}"].value)).quantize(Decimal("0.01"))
    excel_expenses = Decimal(str(worksheet[f"N{summary_row}"].value)).quantize(Decimal("0.01"))
    excel_savings = Decimal(str(worksheet[f"P{summary_row}"].value)).quantize(Decimal("0.01"))

    if total_income != excel_income or total != excel_expenses or actual_savings != excel_savings:
        raise ValueError("Calculated totals do not match Excel")
    

def main():
    file_path = Path(__file__).resolve().parent.parent / "data" / "Spend analyser.xlsx"

    expenses = import_transactions(file_path, "Expense2026", 3, 33)
    incomes = import_income(file_path, "Expense2026", 3, 33)

    total = calculate_total(expenses)   
    total_income = calculate_total(incomes) 
    actual_savings = total_income - total 

    reconcile_totals(file_path, "Expense2026", 34, total_income, total, actual_savings)

    print(f"January income: ${total_income}")   
    print(f"January expenses: ${total}")
    print(f"January savings: ${actual_savings}")
    print("January reconciliation: PASSED")


if __name__ == "__main__":
    main()